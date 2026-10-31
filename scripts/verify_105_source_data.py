# scripts/verify_105_source_data.py
import openpyxl, json, re, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = r'C:\Users\truongg\OneDrive - Thuyloi University\Tài liệu\01. khoi nghiep in 3d\Top_Mau_In_3D_Doc_La_Xu_The_MakerWorld.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)

# 1. Read category sheets
cat_sheet_names = [
    'Trang_Diem_Vanity',
    'Decor_Den_Noi_That',
    'Hoa_Decor_Kim_Loai',
    'Do_Choi_Moc_Khoa',
    'Moc_Khoa_Clicker_Viral',
    'Hop_De_Ban_Tien_Ich',
    'Ca_Nhan_Hoa_NFC'
]

sheet_details = {}
for sname in cat_sheet_names:
    sheet = wb[sname]
    sheet_details[sname] = {}
    for r in range(3, sheet.max_row + 1):
        stt = sheet.cell(r, 1).value
        name_cell = sheet.cell(r, 3).value
        author_cell = sheet.cell(r, 4).value
        score_cell = sheet.cell(r, 5).value
        analysis_cell = sheet.cell(r, 6).value
        mech_cell = sheet.cell(r, 7).value
        benefit_cell = sheet.cell(r, 8).value
        stats_cell = sheet.cell(r, 9).value
        video_cell = sheet.cell(r, 10).value
        mat_cell = sheet.cell(r, 11).value
        weight_cell = sheet.cell(r, 12).value
        time_cell = sheet.cell(r, 13).value
        price_cell = sheet.cell(r, 15).value
        license_cell = sheet.cell(r, 18).value
        
        if not name_cell:
            continue
        m = re.search(r'#(\d+)', str(name_cell))
        mid = m.group(1) if m else None
        key = mid or str(stt)
        
        sheet_details[sname][key] = {
            'stt': stt,
            'name': str(name_cell).strip(),
            'author_orig': str(author_cell).strip() if author_cell else '',
            'score': str(score_cell).strip() if score_cell else '',
            'analysis': str(analysis_cell).strip() if analysis_cell else '',
            'mech': str(mech_cell).strip() if mech_cell else '',
            'benefit': str(benefit_cell).strip() if benefit_cell else '',
            'video': str(video_cell).strip() if video_cell else '',
            'material': str(mat_cell).strip() if mat_cell else '',
            'weight': weight_cell,
            'print_time': time_cell,
            'price': int(price_cell) if isinstance(price_cell, (int, float)) and price_cell > 0 else None,
            'license': str(license_cell).strip() if license_cell else ''
        }

# 2. Read Du_lieu_nguon
sheet_src = wb['Du_lieu_nguon']
items_by_cat = {c: [] for c in cat_sheet_names}

for r in range(3, sheet_src.max_row + 1):
    mid = str(sheet_src.cell(r, 1).value or '').strip()
    cat = str(sheet_src.cell(r, 2).value or '').strip()
    orig_name = str(sheet_src.cell(r, 3).value or '').strip()
    vn_name = str(sheet_src.cell(r, 4).value or '').strip()
    author = str(sheet_src.cell(r, 5).value or '').strip()
    url = str(sheet_src.cell(r, 6).value or '').strip()
    downloads = int(sheet_src.cell(r, 8).value or 0)
    likes = int(sheet_src.cell(r, 9).value or 0)
    prints = int(sheet_src.cell(r, 10).value or 0)
    weight = sheet_src.cell(r, 11).value or 0
    time_s = sheet_src.cell(r, 12).value or 0
    a_score = float(sheet_src.cell(r, 13).value or 5)
    u_score = float(sheet_src.cell(r, 14).value or 5)
    f_score = float(sheet_src.cell(r, 15).value or 5)
    sel_score = (100.0 * (a_score + u_score + f_score)) / 15.0
    reason = str(sheet_src.cell(r, 17).value or '').strip()
    license_name = str(sheet_src.cell(r, 18).value or '').strip()
    commercial = str(sheet_src.cell(r, 19).value or '').strip()
    
    if not mid or cat not in items_by_cat:
        continue
        
    detail = sheet_details.get(cat, {}).get(mid, {})
    price = detail.get('price')
    
    item = {
        'mid': mid,
        'cat': cat,
        'vn_name': vn_name,
        'orig_name': orig_name,
        'author': author,
        'url': url,
        'downloads': downloads,
        'likes': likes,
        'prints': prints,
        'weight': weight,
        'time_s': time_s,
        'a_score': a_score,
        'u_score': u_score,
        'f_score': f_score,
        'sel_score': sel_score,
        'reason': reason,
        'license_name': license_name,
        'commercial': commercial,
        'price': price,
        'detail': detail
    }
    items_by_cat[cat].append(item)

print("=== CATEGORY COUNTS ===")
for cat, items in items_by_cat.items():
    print(f"{cat:25}: {len(items)} items")

# Check Hero selection for each category
print("\n=== HERO CANDIDATES (Codex Ordering) ===")
for cat, items in items_by_cat.items():
    # Sort key: (-sel_score, -prints, -downloads, -likes, int(mid))
    sorted_items = sorted(items, key=lambda x: (-x['sel_score'], -x['prints'], -x['downloads'], -x['likes'], int(x['mid'])))
    hero = sorted_items[0]
    print(f"[{cat}] Hero: mw-cn-{hero['mid']} (Score: {hero['sel_score']:.1f}, Prints: {hero['prints']}, Downloads: {hero['downloads']}, Likes: {hero['likes']}) -> {hero['vn_name'][:50]}")
