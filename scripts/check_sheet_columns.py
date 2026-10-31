import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = r'C:\Users\truongg\OneDrive - Thuyloi University\Tài liệu\01. khoi nghiep in 3d\Top_Mau_In_3D_Doc_La_Xu_The_MakerWorld.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
sheet = wb['Trang_Diem_Vanity']
headers = [str(sheet.cell(2, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
for idx, h in enumerate(headers, 1):
    print(f"Col {idx}: {h}")
