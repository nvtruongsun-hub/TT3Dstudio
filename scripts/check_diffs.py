import openpyxl, json, sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

f = r'C:\Users\truongg\.gemini\antigravity\brain\9324bcca-30cb-4b20-b89a-f522610fa9cc\.user_uploaded\media_1792552368163_d3adc6ce.xlsx'
wb = openpyxl.load_workbook(f, data_only=True)

with open('products.json', 'r', encoding='utf-8') as pf:
    products = json.load(pf)

differences = []

for sname in ['Trang_Diem_Vanity', 'Decor_Den_Noi_That', 'Do_Choi_Moc_Khoa', 'Hop_De_Ban_Tien_Ich', 'Ca_Nhan_Hoa_NFC']:
    sheet = wb[sname]
    headers = [str(sheet.cell(2, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
    for r in range(3, 14):
        stt = sheet.cell(r, 1).value
        name = str(sheet.cell(r, 3).value or '').strip()
        m = re.search(r'#(\d+)', name)
        mid = m.group(1) if m else None
        
        p = next((x for x in products if mid and (mid in x['id'] or mid in x.get('makerworldUrl', ''))), None)
        if not p:
            differences.append(f"Missing in JSON: {mid} - {name[:40]}")
            continue
            
        for c, h in enumerate(headers, 1):
            if 'giá bán lẻ' in h.lower():
                val = sheet.cell(r, c).value
                if val is not None and int(val) != p.get('price'):
                    differences.append(f"Price diff {p['id']}: Excel={val} vs JSON={p.get('price')}")
            if 'tác giả' in h.lower():
                author_val = str(sheet.cell(r, c).value or '').strip()
                # check author

print(f"Total differences found: {len(differences)}")
for d in differences:
    print(" -", d)
