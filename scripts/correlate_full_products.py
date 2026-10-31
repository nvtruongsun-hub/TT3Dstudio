import openpyxl, json, sys, io, re

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = r'C:\Users\truongg\OneDrive - Thuyloi University\Tài liệu\01. khoi nghiep in 3d\Top_Mau_In_3D_Doc_La_Xu_The_MakerWorld.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)

# 1. Read Tong_Hop_Kinh_Te for prices
sheet_econ = wb['Tong_Hop_Kinh_Te']
econ_prices = {}
for r in range(24, sheet_econ.max_row + 1):
    # Col 1: STT, Col 2: Phân khúc, Col 3: Tên, Col 4: Nhựa (g), Col 5: Giờ In, Col 11: Giá Bán Lẻ Đề Xuất
    stt = sheet_econ.cell(r, 1).value
    cat = sheet_econ.cell(r, 2).value
    name = sheet_econ.cell(r, 3).value
    weight = sheet_econ.cell(r, 4).value
    hours = sheet_econ.cell(r, 5).value
    price = sheet_econ.cell(r, 11).value
    if name and price:
        # try find model id in name
        m = re.search(r'#(\d+)', str(name))
        mid = m.group(1) if m else None
        key = (cat, stt)
        econ_prices[key] = {
            'price': int(price) if isinstance(price, (int, float)) else None,
            'weight': weight,
            'hours': hours,
            'model_id': mid,
            'name': str(name).strip()
        }

# 2. Read each category sheet for row-by-row details
sheet_details = {}
cat_sheet_names = [
    'Trang_Diem_Vanity',
    'Decor_Den_Noi_That',
    'Hoa_Decor_Kim_Loai',
    'Do_Choi_Moc_Khoa',
    'Moc_Khoa_Clicker_Viral',
    'Hop_De_Ban_Tien_Ich',
    'Ca_Nhan_Hoa_NFC'
]

for sname in cat_sheet_names:
    sheet = wb[sname]
    sheet_details[sname] = {}
    for r in range(3, sheet.max_row + 1):
        stt = sheet.cell(r, 1).value
        name_cell = sheet.cell(r, 3).value
        author_cell = sheet.cell(r, 4).value
        score_cell = sheet.cell(r, 5).value
        analysis_cell = sheet.cell(r, 6).value
        mech_cell = sheet.cell(r, 7).value
        benefit_cell = sheet.cell(r, 8).value
        stats_cell = sheet.cell(r, 9).value
        video_cell = sheet.cell(r, 10).value
        mat_cell = sheet.cell(r, 11).value
        weight_cell = sheet.cell(r, 12).value
        time_cell = sheet.cell(r, 13).value
        price_cell = sheet.cell(r, 15).value
        license_cell = sheet.cell(r, 18).value
        
        if not name_cell:
            continue
        m = re.search(r'#(\d+)', str(name_cell))
        mid = m.group(1) if m else None
        
        sheet_details[sname][mid or str(stt)] = {
            'stt': stt,
            'name': str(name_cell).strip(),
            'author_orig': str(author_cell).strip() if author_cell else '',
            'score': str(score_cell).strip() if score_cell else '',
            'analysis': str(analysis_cell).strip() if analysis_cell else '',
            'benefit': str(benefit_cell).strip() if benefit_cell else '',
            'video': str(video_cell).strip() if video_cell else '',
            'material': str(mat_cell).strip() if mat_cell else '',
            'weight': weight_cell,
            'print_time': time_cell,
            'price': int(price_cell) if isinstance(price_cell, (int, float)) and price_cell > 0 else None,
            'license': str(license_cell).strip() if license_cell else ''
        }

# 3. Read Du_lieu_nguon
sheet_src = wb['Du_lieu_nguon']
full_products = []

for r in range(3, sheet_src.max_row + 1):
    mid = str(sheet_src.cell(r, 1).value or '').strip()
    cat = str(sheet_src.cell(r, 2).value or '').strip()
    orig_name = str(sheet_src.cell(r, 3).value or '').strip()
    vn_name = str(sheet_src.cell(r, 4).value or '').strip()
    author = str(sheet_src.cell(r, 5).value or '').strip()
    url = str(sheet_src.cell(r, 6).value or '').strip()
    downloads = sheet_src.cell(r, 8).value or 0
    likes = sheet_src.cell(r, 9).value or 0
    prints = sheet_src.cell(r, 10).value or 0
    weight = sheet_src.cell(r, 11).value or 0
    time_s = sheet_src.cell(r, 12).value or 0
    a_score = sheet_src.cell(r, 13).value or 5
    u_score = sheet_src.cell(r, 14).value or 5
    f_score = sheet_src.cell(r, 15).value or 5
    sel_score = sheet_src.cell(r, 16).value or 100
    reason = str(sheet_src.cell(r, 17).value or '').strip()
    license_name = str(sheet_src.cell(r, 18).value or '').strip()
    commercial = str(sheet_src.cell(r, 19).value or '').strip()
    thumb = str(sheet_src.cell(r, 21).value or '').strip()
    
    if not mid:
        continue
        
    # find price in sheet_details
    detail = sheet_details.get(cat, {}).get(mid, {})
    price = detail.get('price')
    
    full_products.append({
        'mid': mid,
        'cat': cat,
        'vn_name': vn_name,
        'orig_name': orig_name,
        'author': author,
        'price': price,
        'detail': detail
    })

print(f"Total processed products: {len(full_products)}")
missing_price = [p for p in full_products if not p['price']]
print(f"Products missing price: {len(missing_price)}")
for mp in missing_price[:5]:
    print(f" * Missing price: {mp['mid']} in {mp['cat']}: {mp['vn_name'][:50]}")
