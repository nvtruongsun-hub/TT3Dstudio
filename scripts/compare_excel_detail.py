import openpyxl, json, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

f = r'C:\Users\truongg\.gemini\antigravity\brain\9324bcca-30cb-4b20-b89a-f522610fa9cc\.user_uploaded\media_1792552368163_d3adc6ce.xlsx'
wb = openpyxl.load_workbook(f, data_only=True)

with open('products.json', 'r', encoding='utf-8') as pf:
    products = json.load(pf)

p_map = {p['id']: p for p in products}

print("=== Comparing Excel vs products.json ===")
for sname in ['Trang_Diem_Vanity', 'Decor_Den_Noi_That', 'Do_Choi_Moc_Khoa', 'Hop_De_Ban_Tien_Ich', 'Ca_Nhan_Hoa_NFC']:
    sheet = wb[sname]
    print(f"\n--- Sheet: {sname} ---")
    headers = [str(sheet.cell(2, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
    print("Headers:", [h for h in headers if h][:8])
    for r in range(3, sheet.max_row + 1):
        stt = sheet.cell(r, 1).value
        if not stt or not str(stt).strip().isdigit():
            continue
        name_val = str(sheet.cell(r, 3).value or '').strip()
        author_val = str(sheet.cell(r, 4).value or '').strip()
        score_val = str(sheet.cell(r, 5).value or '').strip()
        analysis_val = str(sheet.cell(r, 6).value or '').strip()
        # let's find matching product in products.json
        # find by model id
        import re
        m = re.search(r'#(\d+)', name_val)
        mid = m.group(1) if m else None
        
        # let's see which product has this mid
        matched_p = None
        for p in products:
            if mid and (mid in p['id'] or mid in p.get('makerworldUrl', '') or mid in p.get('title', '')):
                matched_p = p
                break
        
        if matched_p:
            # print differences in price, or other attributes
            # what columns are in the sheet?
            cols = {headers[c-1]: sheet.cell(r, c).value for c in range(1, sheet.max_column + 1) if c-1 < len(headers) and headers[c-1]}
            # let's see pricing or specs in cols
            # print some key columns
            price_keys = [k for k in cols.keys() if 'giá' in k.lower() or 'chi phí' in k.lower() or 'bán' in k.lower() or 'vnd' in k.lower()]
            print(f"Row {r} (STT {stt}) [ID #{mid}]: Json Price = {matched_p.get('price')}, Json Title = '{matched_p.get('title')}'")
            if price_keys:
                print(f"   Excel Price info: { {k: cols[k] for k in price_keys} }")
