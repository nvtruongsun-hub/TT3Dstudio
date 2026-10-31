import openpyxl, json, sys, io, re, os

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = r'C:\Users\truongg\OneDrive - Thuyloi University\Tài liệu\01. khoi nghiep in 3d\Top_Mau_In_3D_Doc_La_Xu_The_MakerWorld.xlsx'
print("Loading workbook (this might take a few seconds because file is 64MB)...")
wb = openpyxl.load_workbook(fpath, data_only=True)

with open('products.json', 'r', encoding='utf-8') as pf:
    current_products = json.load(pf)

current_ids = {p['id']: p for p in current_products}
current_model_ids = {p.get('modelId'): p for p in current_products if p.get('modelId')}

print(f"Current products.json has {len(current_products)} items.")

# Inspect Du_lieu_nguon
sheet_src = wb['Du_lieu_nguon']
print(f"\n=== Sheet Du_lieu_nguon (rows={sheet_src.max_row}, cols={sheet_src.max_column}) ===")
headers = [str(sheet_src.cell(2, c).value or '').strip() for c in range(1, sheet_src.max_column + 1)]
print("Headers in Du_lieu_nguon:", headers)

source_items = []
for r in range(3, sheet_src.max_row + 1):
    mid = sheet_src.cell(r, 1).value
    cat = sheet_src.cell(r, 2).value
    orig_name = sheet_src.cell(r, 3).value
    vn_name = sheet_src.cell(r, 4).value
    author = sheet_src.cell(r, 5).value
    url = sheet_src.cell(r, 6).value
    if mid or vn_name:
        source_items.append({
            'row': r,
            'model_id': str(mid).strip() if mid else None,
            'category': str(cat).strip() if cat else None,
            'orig_name': str(orig_name).strip() if orig_name else '',
            'vn_name': str(vn_name).strip() if vn_name else '',
            'author': str(author).strip() if author else '',
            'url': str(url).strip() if url else ''
        })

print(f"Total items in Du_lieu_nguon: {len(source_items)}")

# Count by category
cat_counts = {}
for item in source_items:
    c = item['category'] or 'Unknown'
    cat_counts[c] = cat_counts.get(c, 0) + 1

print("\nItem counts by category in Du_lieu_nguon:")
for c, cnt in sorted(cat_counts.items()):
    print(f" - {c}: {cnt} items")

# Check overlap with current_products
existing_count = 0
new_items = []
for item in source_items:
    mid = item['model_id']
    if mid in current_model_ids or any(mid and mid in p['id'] for p in current_products):
        existing_count += 1
    else:
        new_items.append(item)

print(f"\nExisting items matched: {existing_count}")
print(f"Brand new items to add: {len(new_items)}")
print(f"Sample of new items (first 10):")
for item in new_items[:10]:
    print(f" * [Row {item['row']} | {item['category']} | ID #{item['model_id']}]: {item['vn_name'][:70]} (Author: {item['author']})")
