import openpyxl, json, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = r'C:\Users\truongg\OneDrive - Thuyloi University\Tài liệu\01. khoi nghiep in 3d\Top_Mau_In_3D_Doc_La_Xu_The_MakerWorld.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)

# Inspect Tong_Hop_Kinh_Te
sheet_econ = wb['Tong_Hop_Kinh_Te']
print(f"=== Sheet Tong_Hop_Kinh_Te (rows={sheet_econ.max_row}, cols={sheet_econ.max_column}) ===")
for r in range(1, 30):
    vals = [str(sheet_econ.cell(r, c).value or '').strip() for c in range(1, sheet_econ.max_column + 1)]
    if any(vals):
        print(f"R{r}:", ' | '.join([v for v in vals if v]))

# Also check headers of a couple of specific sheets: Hoa_Decor_Kim_Loai and Moc_Khoa_Clicker_Viral
for sname in ['Hoa_Decor_Kim_Loai', 'Moc_Khoa_Clicker_Viral']:
    sheet = wb[sname]
    print(f"\n=== Sheet {sname} (rows={sheet.max_row}) ===")
    headers = [str(sheet.cell(2, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
    print("Headers:", [h for h in headers if h][:8])
    for r in range(3, min(sheet.max_row + 1, 6)):
        vals = [str(sheet.cell(r, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
        print(f"  R{r}:", ' | '.join([v for v in vals if v][:5]))
